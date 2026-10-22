"""Read Microsoft Forms Excel exports (or a simple CSV)."""

from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook

FIRST_RE = re.compile(r"^(?P<q>.+?)\s+[—–-]\s+first answer\s*$", re.I)
SECOND_RE = re.compile(r"^(?P<q>.+?)\s+[—–-]\s+second answer(?:\s+\(optional\))?\s*$", re.I)
IDENTITY = {"id", "start time", "completion time", "email", "name", "last modified time"}


def parse_survey(path: str | Path) -> tuple[list[str], dict[str, list[tuple[str, list[str]]]]]:
    """Return (question_order, {question: [(respondent, [answers...]), ...]})."""
    p = Path(path)
    if p.suffix.lower() in {".csv"}:
        rows, headers = _read_csv(p)
    else:
        rows, headers = _read_xlsx(p)
    return _group(headers, rows)


def _read_xlsx(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    body = [list(r) for r in rows[1:] if any(x is not None and str(x).strip() for x in r)]
    return body, headers


def _read_csv(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return [], []
    headers = [c.strip() for c in rows[0]]
    body = rows[1:]
    return body, headers


def _group(headers: list[str], rows: list[list]):
    name_i = _find_header(headers, "name")
    email_i = _find_header(headers, "email")
    id_i = _find_header(headers, "id")
    first_map: dict[str, int] = {}
    second_map: dict[str, int] = {}
    order: list[str] = []
    for i, h in enumerate(headers):
        m1 = FIRST_RE.match(h.strip())
        m2 = SECOND_RE.match(h.strip())
        if m1:
            q = m1.group("q").strip()
            first_map[q] = i
            if q not in order:
                order.append(q)
        elif m2:
            q = m2.group("q").strip()
            second_map[q] = i
            if q not in order:
                order.append(q)

    if not order:
        raise ValueError(
            "No Forms question columns found. Expected headers like "
            "'Name a thing — first answer' and optional '— second answer (optional)'."
        )

    by_q: dict[str, list[tuple[str, list[str]]]] = {q: [] for q in order}
    for r in rows:
        rid = _respondent_id(r, name_i, email_i, id_i)
        for q in order:
            answers = []
            fi = first_map.get(q)
            si = second_map.get(q)
            if fi is not None:
                answers.append(_cell(r, fi))
            if si is not None:
                answers.append(_cell(r, si))
            answers = [a for a in answers if a]
            if answers:
                by_q[q].append((rid, answers))
    return order, by_q


def _find_header(headers: list[str], name: str) -> int | None:
    for i, h in enumerate(headers):
        if h.strip().lower() == name:
            return i
    return None


def _respondent_id(row, name_i, email_i, id_i) -> str:
    for i in (name_i, email_i, id_i):
        if i is None:
            continue
        v = _cell(row, i)
        if v:
            return v
    return "anon"


def _cell(row, i: int) -> str:
    if i >= len(row) or row[i] is None:
        return ""
    return str(row[i]).strip()
